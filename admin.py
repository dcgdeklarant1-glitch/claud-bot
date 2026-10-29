"""Admin panel: mijozlar bazasi va kataloglar. Faqat ADMIN_ID dagi foydalanuvchilar ko'radi."""
import datetime as dt
import io
import json
import logging
from html import escape

from aiogram import Bot, F, Router
from aiogram.filters import Command
from aiogram.types import (Message, CallbackQuery, InlineKeyboardMarkup, InlineKeyboardButton,
                           BotCommand, BotCommandScopeChat, BufferedInputFile)

import certificate
import db
import logistics
import trade
from logistics import admin_ids
from texts import T

PAGE = 10
CATALOGS = {
    "users": "👥 Mijozlar",
    "trades": "🌐 Import / eksport zakazlari",
    "orders": "🚚 Logistika buyurtmalari",
    "certs": "📋 Sertifikat arizalari",
    "calcs": "🧮 Rastamojka hisob-kitoblari",
}
LANG_FLAG = {"uz": "🇺🇿", "ru": "🇷🇺", "en": "🇬🇧", "zh": "🇨🇳"}


def is_admin(event) -> bool:
    return event.from_user is not None and event.from_user.id in admin_ids()


# ---------- Yordamchilar ----------
def full_time(iso: str | None) -> str:
    """2026-10-29T14:30:00+05:00 -> 2026-10-29 14:30 (Excel uchun)"""
    return iso[:16].replace("T", " ") if iso else ""

def short_time(iso: str | None) -> str:
    """2026-10-29T14:30:00+05:00 -> 29.10 14:30"""
    return f"{iso[8:10]}.{iso[5:7]} {iso[11:16]}" if iso else "—"

def who(row: dict) -> str:
    name = escape(row.get("name") or str(row.get("user_id") or row.get("id")))
    uid = row.get("user_id") or row.get("id")
    text = f'<a href="tg://user?id={uid}">{name}</a>'
    if row.get("username"):
        text += f" @{escape(row['username'])}"
    return text

def place_plain(place: dict) -> str:
    if not place:
        return ""
    if "lat" in place:
        return f"https://maps.google.com/?q={place['lat']},{place['lon']}"
    return place["text"]

def option_plain(d: dict, field: str) -> str:
    if d.get(field + "_text"):
        return d[field + "_text"]
    return T["uz"]["lg_" + field].get(d.get(field), d.get(field) or "")

def menu_kb():
    rows = [[InlineKeyboardButton(text=title, callback_data=f"adm:list:{cat}:0")]
            for cat, title in CATALOGS.items()]
    rows.append([InlineKeyboardButton(text="📥 Excel'ga yuklab olish", callback_data="adm:xlsx",
                                      style="success")])
    rows.append([InlineKeyboardButton(text="🔄 Yangilash", callback_data="adm:menu")])
    return InlineKeyboardMarkup(inline_keyboard=rows)

def page_kb(cat: str, page: int, total: int):
    nav = []
    if page > 0:
        nav.append(InlineKeyboardButton(text="⬅️", callback_data=f"adm:list:{cat}:{page - 1}"))
    pages = max(1, (total + PAGE - 1) // PAGE)
    nav.append(InlineKeyboardButton(text=f"{page + 1} / {pages}", callback_data="adm:noop"))
    if (page + 1) * PAGE < total:
        nav.append(InlineKeyboardButton(text="➡️", callback_data=f"adm:list:{cat}:{page + 1}"))
    return InlineKeyboardMarkup(inline_keyboard=[
        nav,
        [InlineKeyboardButton(text="🏠 Admin panel", callback_data="adm:menu")],
    ])


# ---------- Statistika ----------
async def stats_text() -> str:
    today = db.today()
    week_ago = (dt.datetime.now(db.TZ) - dt.timedelta(days=7)).isoformat(timespec="seconds")
    q = db.fetchval
    users = await q("SELECT COUNT(*) FROM users")
    users_new = await q("SELECT COUNT(*) FROM users WHERE first_seen >= $1", today)
    users_active = await q("SELECT COUNT(*) FROM users WHERE last_seen >= $1", week_ago)
    trades = await q("SELECT COUNT(*) FROM trades")
    trades_new = await q("SELECT COUNT(*) FROM trades WHERE created >= $1", today)
    orders = await q("SELECT COUNT(*) FROM orders")
    orders_new = await q("SELECT COUNT(*) FROM orders WHERE created >= $1", today)
    certs = await q("SELECT COUNT(*) FROM certs")
    certs_wait = await q("SELECT COUNT(*) FROM certs WHERE answer IS NULL")
    calcs = await q("SELECT COUNT(*) FROM calcs")
    calcs_new = await q("SELECT COUNT(*) FROM calcs WHERE created >= $1", today)
    return (
        "🛠 <b>Admin panel</b>\n\n"
        f"👥 Mijozlar: <b>{users}</b> (bugun yangi: {users_new}, 7 kunda faol: {users_active})\n"
        f"🌐 Import / eksport zakazlari: <b>{trades}</b> (bugun: {trades_new})\n"
        f"🚚 Logistika buyurtmalari: <b>{orders}</b> (bugun: {orders_new})\n"
        f"📋 Sertifikat arizalari: <b>{certs}</b> (⏳ javob kutmoqda: {certs_wait})\n"
        f"🧮 Hisob-kitoblar: <b>{calcs}</b> (bugun: {calcs_new})\n\n"
        "Katalogni tanlang 👇"
    )


# ---------- Kataloglar ro'yxati ----------
async def list_text(cat: str, page: int) -> tuple[str, int]:
    off = page * PAGE
    if cat == "users":
        total = await db.fetchval("SELECT COUNT(*) FROM users")
        rows = await db.fetch("SELECT * FROM users ORDER BY last_seen DESC LIMIT $1 OFFSET $2", PAGE, off)
        lines = [f"{LANG_FLAG.get(r['lang'], '🌐')} {who(r)}"
                 f"{' · 📱 ' + escape(r['phone']) if r['phone'] else ''}\n"
                 f"    🕒 {short_time(r['last_seen'])} · /m{r['id']}" for r in rows]
    elif cat == "trades":
        total = await db.fetchval("SELECT COUNT(*) FROM trades")
        rows = await db.fetch(
            "SELECT t.*, u.name, u.username FROM trades t LEFT JOIN users u ON u.id = t.user_id "
            "ORDER BY t.no DESC LIMIT $1 OFFSET $2", PAGE, off)
        lines = [f"<b>#{r['no']}</b> · {short_time(r['created'])} · {who(r)}\n"
                 f"    {T['uz']['tr_direction'].get(r['direction'], '')} · 📦 {escape((r['product'] or '')[:30])}"
                 f" · 📱 {escape(r['phone'])} · /z{r['no']}" for r in rows]
    elif cat == "orders":
        total = await db.fetchval("SELECT COUNT(*) FROM orders")
        rows = await db.fetch(
            "SELECT o.*, u.name, u.username FROM orders o LEFT JOIN users u ON u.id = o.user_id "
            "ORDER BY o.no DESC LIMIT $1 OFFSET $2", PAGE, off)
        lines = []
        for r in rows:
            d = json.loads(r["data"])
            route = " → ".join("📍 lokatsiya" if "lat" in d[f] else escape(d[f]["text"][:25])
                               for f in ("origin", "destination"))
            lines.append(f"<b>#{r['no']}</b> · {short_time(r['created'])} · {who(r)}\n"
                         f"    {route} · /b{r['no']}")
    elif cat == "certs":
        total = await db.fetchval("SELECT COUNT(*) FROM certs")
        rows = await db.fetch(
            "SELECT c.*, u.name, u.username FROM certs c LEFT JOIN users u ON u.id = c.user_id "
            "ORDER BY c.no DESC LIMIT $1 OFFSET $2", PAGE, off)
        lines = [f"{'✅' if r['answer'] else '⏳'} <b>#{r['no']}</b> · {short_time(r['created'])} · {who(r)}\n"
                 f"    📦 {escape((r['product'] or '')[:40])} · /s{r['no']}" for r in rows]
    else:
        total = await db.fetchval("SELECT COUNT(*) FROM calcs")
        rows = await db.fetch(
            "SELECT c.*, u.name, u.username FROM calcs c LEFT JOIN users u ON u.id = c.user_id "
            "ORDER BY c.id DESC LIMIT $1 OFFSET $2", PAGE, off)
        lines = [f"{short_time(r['created'])} · {who(r)}\n"
                 f"    {T['uz']['fuel_' + r['fuel']]}, {r['year']}"
                 f"{', ' + str(r['cc']) + ' sm³' if r['cc'] else ''} · ${r['price']:,.0f} → "
                 f"rastamojka ${r['total_usd']:,.0f} · /m{r['user_id']}" for r in rows]
    body = "\n\n".join(lines) if lines else "<i>Hozircha bo'sh.</i>"
    return f"<b>{CATALOGS[cat]}</b> (jami: {total})\n\n{body}", total


# ---------- Kartochkalar ----------
async def user_card(uid: int) -> str | None:
    rows = await db.fetch("SELECT * FROM users WHERE id = $1", uid)
    if not rows:
        return None
    u = rows[0]
    trades = await db.fetch("SELECT no, created FROM trades WHERE user_id = $1 ORDER BY no DESC", uid)
    orders = await db.fetch("SELECT no, created FROM orders WHERE user_id = $1 ORDER BY no DESC", uid)
    certs = await db.fetch("SELECT no, created, answer FROM certs WHERE user_id = $1 ORDER BY no DESC", uid)
    calcs = await db.fetchval("SELECT COUNT(*) FROM calcs WHERE user_id = $1", uid)
    text = (f"👤 <b>Mijoz</b>\n\n"
            f"Ism: {who(u)}\n"
            f"ID: <code>{u['id']}</code>\n"
            f"📱 Telefon: {escape(u['phone']) if u['phone'] else '—'}\n"
            f"🌐 Til: {LANG_FLAG.get(u['lang'], '')} {u['lang'] or '—'}\n"
            f"🆕 Birinchi kirgan: {short_time(u['first_seen'])}\n"
            f"🕒 Oxirgi faollik: {short_time(u['last_seen'])}\n\n"
            f"🌐 Import / eksport zakazlari: {len(trades)}\n")
    text += "".join(f"    #{z['no']} · {short_time(z['created'])} · /z{z['no']}\n" for z in trades[:5])
    text += f"🚚 Logistika buyurtmalari: {len(orders)}\n"
    text += "".join(f"    #{o['no']} · {short_time(o['created'])} · /b{o['no']}\n" for o in orders[:5])
    text += f"📋 Sertifikat arizalari: {len(certs)}\n"
    text += "".join(f"    {'✅' if c['answer'] else '⏳'} #{c['no']} · {short_time(c['created'])} · /s{c['no']}\n"
                    for c in certs[:5])
    text += f"🧮 Hisob-kitoblar: {calcs}"
    return text


# ---------- Excel ----------
async def build_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = []

    users = await db.fetch("SELECT * FROM users ORDER BY first_seen")
    sheets.append(("Mijozlar",
                   ["ID", "Ism", "Username", "Telefon", "Til", "Birinchi kirgan", "Oxirgi faollik"],
                   [[u["id"], u["name"], u["username"], u["phone"], u["lang"], full_time(u["first_seen"]),
                     full_time(u["last_seen"])]
                    for u in users]))

    rows = []
    for r in await db.fetch("SELECT t.*, u.name FROM trades t LEFT JOIN users u ON u.id = t.user_id ORDER BY t.no"):
        d = json.loads(r["data"])
        rows.append([r["no"], full_time(r["created"]), r["user_id"], r["name"], r["phone"],
                     T["uz"]["tr_direction"].get(d.get("direction"), ""), d.get("product"),
                     d.get("country"), d.get("volume"), d.get("name")])
    sheets.append(("Import-eksport",
                   ["№", "Sana", "Mijoz ID", "Ism (Telegram)", "Telefon", "Xizmat", "Tovar", "Davlat",
                    "Hajmi", "Ism / kompaniya"], rows))

    rows = []
    for r in await db.fetch("SELECT o.*, u.name FROM orders o LEFT JOIN users u ON u.id = o.user_id ORDER BY o.no"):
        d = json.loads(r["data"])
        rows.append([r["no"], full_time(r["created"]), r["user_id"], r["name"], r["phone"],
                     place_plain(d.get("origin")), place_plain(d.get("destination")),
                     option_plain(d, "cargo"), option_plain(d, "weight"),
                     option_plain(d, "volume"), option_plain(d, "ready")])
    sheets.append(("Logistika",
                   ["№", "Sana", "Mijoz ID", "Ism", "Telefon", "Qayerdan", "Qayerga",
                    "Yuk turi", "Og'irligi", "Hajmi", "Tayyor bo'ladi"], rows))

    rows = []
    for r in await db.fetch("SELECT c.*, u.name FROM certs c LEFT JOIN users u ON u.id = c.user_id ORDER BY c.no"):
        d = json.loads(r["data"])
        rows.append([r["no"], full_time(r["created"]), r["user_id"], r["name"], r["phone"], d.get("product"),
                     d.get("code"), d.get("country"), T["uz"]["ct_purpose"].get(d.get("purpose"), ""),
                     "ha" if d.get("file") else "yo'q", r["answer"], full_time(r["answered_at"])])
    sheets.append(("Sertifikatlar",
                   ["№", "Sana", "Mijoz ID", "Ism", "Telefon", "Tovar", "TIF TN", "Davlat",
                    "Maqsad", "Rasm/hujjat", "Mutaxassis javobi", "Javob sanasi"], rows))

    calcs = await db.fetch("SELECT c.*, u.name FROM calcs c LEFT JOIN users u ON u.id = c.user_id ORDER BY c.id")
    sheets.append(("Hisob-kitoblar",
                   ["Sana", "Mijoz ID", "Ism", "Yoqilg'i", "Yili", "Hajmi (sm³)", "Narxi ($)",
                    "Yetkazish ($)", "Rastamojka (so'm)", "Rastamojka ($)"],
                   [[full_time(c["created"]), c["user_id"], c["name"], T["uz"]["fuel_" + c["fuel"]], c["year"],
                     c["cc"], c["price"], c["delivery"], round(c["total_uzs"]), round(c["total_usd"])]
                    for c in calcs]))

    wb.remove(wb.active)
    for title, header, data in sheets:
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in data:
            ws.append(row)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def set_admin_commands(bot: Bot):
    """Adminlarning buyruqlar menyusida /admin ko'rinadi (boshqalarda ko'rinmaydi)."""
    commands = [BotCommand(command="start", description="Bosh menyu"),
                BotCommand(command="admin", description="🛠 Admin panel: mijozlar bazasi"),
                BotCommand(command="setkurs", description="💱 CBU kursini kiritish"),
                BotCommand(command="setlogo", description="🖼 Bot rasmini logotipga almashtirish")]
    for admin in admin_ids():
        try:
            await bot.set_my_commands(commands, scope=BotCommandScopeChat(chat_id=admin))
        except Exception as e:
            logging.warning("Admin (%s) buyruqlarini o'rnatib bo'lmadi: %r", admin, e)


def build_router() -> Router:
    router = Router(name="admin")

    @router.message(Command("admin"))
    async def admin_panel(m: Message):
        if not is_admin(m):
            return await m.answer(f"⛔️ Bu bo'lim faqat adminlar uchun.\nSizning ID: <code>{m.from_user.id}</code>")
        await m.answer(await stats_text(), reply_markup=menu_kb())

    @router.callback_query(F.data == "adm:menu", is_admin)
    async def menu(c: CallbackQuery):
        await c.answer()
        try:
            await c.message.edit_text(await stats_text(), reply_markup=menu_kb())
        except Exception:
            pass  # "message is not modified" — ma'lumot o'zgarmagan

    @router.callback_query(F.data.regexp(r"^adm:list:(users|trades|orders|certs|calcs):\d+$"), is_admin)
    async def show_list(c: CallbackQuery):
        _, _, cat, page = c.data.split(":")
        text, total = await list_text(cat, int(page))
        await c.answer()
        await c.message.edit_text(text, reply_markup=page_kb(cat, int(page), total),
                                  disable_web_page_preview=True)

    @router.callback_query(F.data == "adm:noop")
    async def noop(c: CallbackQuery):
        await c.answer()

    @router.callback_query(F.data == "adm:xlsx", is_admin)
    async def export(c: CallbackQuery):
        await c.answer("⏳ Tayyorlanmoqda...")
        data = await build_xlsx()
        name = f"dcg_baza_{db.today()}.xlsx"
        await c.message.answer_document(BufferedInputFile(data, filename=name),
                                        caption="📥 Mijozlar bazasi: har bir katalog alohida varaqda")

    @router.message(F.text.regexp(r"^/m(\d+)$"), is_admin)
    async def show_user(m: Message):
        card = await user_card(int(m.text[2:]))
        await m.answer(card or "❗️ Bunday mijoz topilmadi.")

    @router.message(F.text.regexp(r"^/b(\d+)$"), is_admin)
    async def show_order(m: Message):
        no = int(m.text[2:])
        rows = await db.fetch("SELECT o.*, u.name, u.username FROM orders o "
                              "LEFT JOIN users u ON u.id = o.user_id WHERE o.no = $1", no)
        if not rows:
            return await m.answer("❗️ Bunday buyurtma topilmadi.")
        r = rows[0]
        await m.answer(f"🚚 <b>Logistika buyurtmasi #{no}</b> · {short_time(r['created'])}\n\n"
                       f"👤 Mijoz: {who(r)} · /m{r['user_id']}\n\n"
                       f"{logistics.summary('uz', json.loads(r['data']))}",
                       disable_web_page_preview=True)

    @router.message(F.text.regexp(r"^/z(\d+)$"), is_admin)
    async def show_trade(m: Message):
        no = int(m.text[2:])
        rows = await db.fetch("SELECT t.*, u.name, u.username FROM trades t "
                              "LEFT JOIN users u ON u.id = t.user_id WHERE t.no = $1", no)
        if not rows:
            return await m.answer("❗️ Bunday zakaz topilmadi.")
        r = rows[0]
        await m.answer(f"🌐 <b>Import / eksport zakazi #{no}</b> · {short_time(r['created'])}\n\n"
                       f"👤 Mijoz: {who(r)} · /m{r['user_id']}\n\n"
                       f"{trade.summary('uz', json.loads(r['data']))}")

    @router.message(F.text.regexp(r"^/s(\d+)$"), is_admin)
    async def show_cert(m: Message):
        no = int(m.text[2:])
        rows = await db.fetch("SELECT c.*, u.name, u.username FROM certs c "
                              "LEFT JOIN users u ON u.id = c.user_id WHERE c.no = $1", no)
        if not rows:
            return await m.answer("❗️ Bunday ariza topilmadi.")
        r = rows[0]
        d = json.loads(r["data"])
        f = d.get("file")
        if f:
            send = m.answer_photo if f["type"] == "photo" else m.answer_document
            await send(f["id"], caption=f"📎 #{no}")
        status = (f"✅ <b>Javob berilgan</b> ({short_time(r['answered_at'])}):\n{escape(r['answer'])}"
                  if r["answer"] else "⏳ <b>Javob kutmoqda</b>")
        await m.answer(f"📋 <b>Sertifikat arizasi #{no}</b> · {short_time(r['created'])}\n\n"
                       f"👤 Mijoz: {who(r)} · /m{r['user_id']}\n\n"
                       f"{certificate.summary('uz', d)}\n\n{status}",
                       reply_markup=certificate.reply_kb(r["user_id"], no))

    return router
